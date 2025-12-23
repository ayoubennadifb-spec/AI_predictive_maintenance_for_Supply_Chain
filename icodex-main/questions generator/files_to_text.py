from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Iterable, List, Tuple

import PyPDF2
from docx import Document
from pptx import Presentation

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
}


class ExtractionError(Exception):
    pass


def extract_text_from_pdf(pdf_path: Path) -> str:
    content: List[str] = []
    try:
        with pdf_path.open("rb") as stream:
            reader = PyPDF2.PdfReader(stream)
            if not reader.pages:
                raise ExtractionError("PDF contains no pages.")
            for idx, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    content.append(f"--- Page {idx} ---\n{text}")
    except Exception as exc:
        raise ExtractionError(f"Failed to extract PDF '{pdf_path.name}': {exc}") from exc
    return "\n\n".join(content)


def extract_text_from_docx(docx_path: Path) -> str:
    try:
        doc = Document(docx_path)
        content = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(content)
    except Exception as exc:
        raise ExtractionError(f"Failed to extract DOCX '{docx_path.name}': {exc}") from exc


def extract_text_from_pptx(pptx_path: Path) -> str:
    try:
        presentation = Presentation(pptx_path)
        slides: List[str] = []
        for idx, slide in enumerate(presentation.slides, start=1):
            parts: List[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text = shape.text.strip()
                    if text:
                        parts.append(text.replace("\n", " "))
                elif hasattr(shape, "table") and shape.table:
                    table_text: List[str] = []
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            table_text.append(" | ".join(cells))
                    if table_text:
                        parts.append("; ".join(table_text))
            if parts:
                slides.append(f"--- Slide {idx} ---\n" + "\n".join(parts))
        return "\n\n".join(slides)
    except Exception as exc:
        raise ExtractionError(f"Failed to extract PPTX '{pptx_path.name}': {exc}") from exc


def extract_text_for_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    if ext == ".pptx":
        return extract_text_from_pptx(path)
    raise ExtractionError(f"Unsupported file type: {ext}")


def list_source_files(directory: Path, selected: Iterable[Path] | None = None) -> List[Path]:
    if selected:
        resolved = [path for path in selected if path.exists() and path.is_file()]
        if resolved:
            return resolved
    files = [path for path in sorted(directory.iterdir()) if path.suffix.lower() in SUPPORTED_EXTENSIONS]
    return files


def merge_texts(frags: List[Tuple[str, str]]) -> str:
    merged: List[str] = []
    for filename, text in frags:
        merged.append(f"===== {filename} =====")
        merged.append(text.strip())
    return "\n\n".join(part for part in merged if part.strip())


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract and merge text from uploaded materials.")
    parser.add_argument("--source-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("merged_raw.txt"))
    parser.add_argument(
        "--files",
        nargs="*",
        type=Path,
        help="Optional explicit list of files (paths) to extract. Defaults to all supported files in the source dir.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source_dir = args.source_dir
    if not source_dir.exists():
        print(f"Source directory not found: {source_dir}")
        sys.exit(1)

    candidates = list_source_files(source_dir, args.files)
    if not candidates:
        print("No supported files found for text extraction.")
        sys.exit(1)

    fragments: List[Tuple[str, str]] = []
    for path in candidates:
        try:
            print(f"Extracting {path.name}...")
            extracted = extract_text_for_file(path)
            if not extracted.strip():
                print(f"Warning: {path.name} produced no text.")
                continue
            fragments.append((path.name, extracted))
        except ExtractionError as exc:
            print(exc)

    if not fragments:
        print("No extractable text from any file.")
        sys.exit(1)

    merged = merge_texts(fragments)
    write_text(args.output, merged)
    for filename, _ in fragments:
        print(f"  included: {filename}")
    print(f"Merged text written to: {args.output}")


if __name__ == "__main__":
    main()
